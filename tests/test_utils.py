import pytest
import json
import csv
from pathlib import Path
from openpyxl import Workbook
from src.utils import load_transactions


@pytest.fixture
def sample_json_data(tmp_path):
    """Создает тестовый JSON файл с транзакциями"""
    data = [
        {"date": "2024-03-20", "amount": "100.0", "currency": "RUB", "description": "Покупка"},
        {"date": "2024-03-21", "amount": "200.0", "currency": "USD", "description": "Продажа"}
    ]
    file_path = tmp_path / "test_transactions.json"
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(data, f)
    return str(file_path)

@pytest.fixture
def sample_csv_data(tmp_path):
    """Создает тестовый CSV файл с транзакциями"""
    data = [
        {"date": "2024-03-20", "amount": "100.0", "currency": "RUB", "description": "Покупка"},
        {"date": "2024-03-21", "amount": "200.0", "currency": "USD", "description": "Продажа"}
    ]
    file_path = tmp_path / "test_transactions.csv"
    with open(file_path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=data[0].keys(), delimiter=';')
        writer.writeheader()
        writer.writerows(data)
    return str(file_path)

@pytest.fixture
def sample_excel_data(tmp_path):
    """Создает тестовый Excel файл с транзакциями"""
    data = [
        {"date": "2024-03-20", "amount": "100.0", "currency": "RUB", "description": "Покупка"},
        {"date": "2024-03-21", "amount": "200.0", "currency": "USD", "description": "Продажа"}
    ]
    file_path = tmp_path / "test_transactions.xlsx"
    wb = Workbook()
    ws = wb.active
    
    # Записываем заголовки
    for col, header in enumerate(data[0].keys(), 1):
        ws.cell(row=1, column=col, value=header)
    
    # Записываем данные
    for row, transaction in enumerate(data, 2):
        for col, value in enumerate(transaction.values(), 1):
            ws.cell(row=row, column=col, value=value)
    
    wb.save(file_path)
    return str(file_path)

def test_load_transactions_json(sample_json_data):
    """Тест загрузки транзакций из JSON файла"""
    transactions = load_transactions(sample_json_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'
    assert transactions[1]['currency'] == 'USD'

def test_load_transactions_csv(sample_csv_data):
    """Тест загрузки транзакций из CSV файла"""
    transactions = load_transactions(sample_csv_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'
    assert transactions[1]['currency'] == 'USD'

def test_load_transactions_excel(sample_excel_data):
    """Тест загрузки транзакций из Excel файла"""
    transactions = load_transactions(sample_excel_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'
    assert transactions[1]['currency'] == 'USD'

def test_load_transactions_invalid_json(tmp_path):
    """Тест загрузки некорректного JSON файла"""
    file_path = tmp_path / "invalid.json"
    file_path.write_text("invalid json")
    transactions = load_transactions(str(file_path))
    assert transactions == []

def test_load_transactions_nonexistent_file():
    """Тест загрузки несуществующего файла"""
    transactions = load_transactions("nonexistent_file.json")
    assert transactions == []

def test_load_transactions_unsupported_format(tmp_path):
    """Тест загрузки файла неподдерживаемого формата"""
    file_path = tmp_path / "test.txt"
    file_path.write_text("test data")
    transactions = load_transactions(str(file_path))
    assert transactions == [] 