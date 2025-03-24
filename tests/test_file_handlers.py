import pytest
import csv
from pathlib import Path
from src.file_handlers import read_csv_file, read_excel_file, read_transactions

@pytest.fixture
def sample_csv_data(tmp_path):
    """Создает тестовый CSV файл с транзакциями"""
    data = [
        {"date": "2024-03-20", "amount": "100.0", "currency": "RUB", "description": "Покупка"},
        {"date": "2024-03-21", "amount": "200.0", "currency": "USD", "description": "Продажа"}
    ]
    file_path = tmp_path / "test_transactions.csv"
    with open(file_path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=data[0].keys(), delimiter=';')
        writer.writeheader()
        writer.writerows(data)
    return str(file_path)

@pytest.fixture
def sample_excel_data(tmp_path):
    """Создает тестовый Excel файл с транзакциями"""
    from openpyxl import Workbook
    data = [
        {"date": "2024-03-20", "amount": "100.0", "currency": "RUB", "description": "Покупка"},
        {"date": "2024-03-21", "amount": "200.0", "currency": "USD", "description": "Продажа"}
    ]
    file_path = tmp_path / "test_transactions.xlsx"
    wb = Workbook()
    ws = wb.active
    
    # Записываем заголовки
    for col, header in enumerate(data[0].keys(), 1):
        ws.cell(row=1, column=col, value=header)
    
    # Записываем данные
    for row, transaction in enumerate(data, 2):
        for col, value in enumerate(transaction.values(), 1):
            ws.cell(row=row, column=col, value=value)
    
    wb.save(file_path)
    return str(file_path)

def test_read_csv_file(sample_csv_data):
    """Тест чтения CSV файла"""
    transactions = read_csv_file(sample_csv_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'
    assert transactions[1]['currency'] == 'USD'

def test_read_excel_file(sample_excel_data):
    """Тест чтения Excel файла"""
    transactions = read_excel_file(sample_excel_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'
    assert transactions[1]['currency'] == 'USD'

def test_read_transactions_csv(sample_csv_data):
    """Тест функции read_transactions для CSV файла"""
    transactions = read_transactions(sample_csv_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'

def test_read_transactions_excel(sample_excel_data):
    """Тест функции read_transactions для Excel файла"""
    transactions = read_transactions(sample_excel_data)
    assert len(transactions) == 2
    assert transactions[0]['amount'] == '100.0'

def test_read_transactions_invalid_file():
    """Тест обработки несуществующего файла"""
    with pytest.raises(FileNotFoundError):
        read_transactions("nonexistent_file.csv")

def test_read_transactions_unsupported_format(tmp_path):
    """Тест обработки неподдерживаемого формата файла"""
    file_path = tmp_path / "test.txt"
    file_path.write_text("test data")
    with pytest.raises(ValueError):
        read_transactions(str(file_path)) 