import csv
from typing import List, Dict, Any
from pathlib import Path
from openpyxl import load_workbook
from .logger import setup_logger

logger = setup_logger(__name__)

def read_csv_file(file_path: str) -> List[Dict[str, Any]]:
    """
    Читает транзакции из CSV файла.
    
    Args:
        file_path: Путь к CSV файлу
        
    Returns:
        Список словарей с транзакциями
    """
    try:
        logger.info(f"Начало чтения CSV файла: {file_path}")
        transactions = []
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f, delimiter=';')
            for row in reader:
                transactions.append(row)
        logger.info(f"Успешно прочитано {len(transactions)} транзакций из CSV файла")
        return transactions
    except Exception as e:
        logger.error(f"Ошибка при чтении CSV файла: {str(e)}")
        raise

def read_excel_file(file_path: str) -> List[Dict[str, Any]]:
    """
    Читает транзакции из Excel файла.
    
    Args:
        file_path: Путь к Excel файлу
        
    Returns:
        Список словарей с транзакциями
    """
    try:
        logger.info(f"Начало чтения Excel файла: {file_path}")
        transactions = []
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        
        # Получаем заголовки из первой строки
        headers = [cell.value for cell in ws[1]]
        
        # Читаем данные из остальных строк
        for row in ws.iter_rows(min_row=2):
            transaction = {}
            for header, cell in zip(headers, row):
                transaction[header] = cell.value
            transactions.append(transaction)
            
        wb.close()
        logger.info(f"Успешно прочитано {len(transactions)} транзакций из Excel файла")
        return transactions
    except Exception as e:
        logger.error(f"Ошибка при чтении Excel файла: {str(e)}")
        raise

def read_transactions(file_path: str) -> List[Dict[str, Any]]:
    """
    Читает транзакции из файла в зависимости от его расширения.
    
    Args:
        file_path: Путь к файлу с транзакциями
        
    Returns:
        Список словарей с транзакциями
    """
    file_path = Path(file_path)
    if not file_path.exists():
        logger.error(f"Файл не найден: {file_path}")
        raise FileNotFoundError(f"Файл не найден: {file_path}")
        
    if file_path.suffix.lower() == '.csv':
        return read_csv_file(str(file_path))
    elif file_path.suffix.lower() in ['.xlsx', '.xls']:
        return read_excel_file(str(file_path))
    else:
        logger.error(f"Неподдерживаемый формат файла: {file_path.suffix}")
        raise ValueError(f"Неподдерживаемый формат файла: {file_path.suffix}") 